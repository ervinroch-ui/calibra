"""Calibra v5 — Presentaciones de control de calidad basadas en el informe 09/10/2026.

No modifica libros Excel. Genera tablas dinámicas, cartas estadísticas y notas
nativas de PowerPoint; conserva las precauciones sobre límites, cantidad de
muestras y estabilidad del proceso.
"""
from __future__ import annotations

import io
import math
import unicodedata
from collections import defaultdict
from datetime import datetime
from statistics import mean

import numpy as np
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from ppt_report import (rgb, text, shape, fmt, advanced, make_plot, limitation,
                        normalize_record, operational_status, period_label,
                        add_issues, add_metric_charts)

NAVY = '#223D63'
INK = '#22334B'
WHITE = '#FFFFFF'
MUTED = '#677890'
LINE = '#DCE3EC'
PALE = '#F2F5F9'
GREEN = '#178148'
RED = '#C52C23'
YELLOW = '#BF8300'
BLUE = '#5D7DAD'
STATUS = {'VERDE':GREEN, 'AMARILLO':YELLOW, 'ROJO':RED, 'NO EVALUABLE':'#596981'}


def notes(slide, paragraphs):
    """PowerPoint's native presenter notes (visible in Notes View, not the slide)."""
    if isinstance(paragraphs, (tuple, list)):
        paragraphs = '\n\n'.join(str(x).strip() for x in paragraphs if str(x).strip())
    slide.notes_slide.notes_text_frame.text = str(paragraphs)


def titled(prs, category, heading, context='', page=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    text(slide, category.upper(), .48, .20, 12.15, .27, 9.5, BLUE, True)
    text(slide, heading, .48, .61, 12.2, .58, 24, NAVY, True)
    if context:
        text(slide, context[:180], .5, 1.23, 12.18, .25, 9.3, MUTED)
    shape(slide, .45, 7.13, 12.45, .014, LINE, radius=False)
    text(slide, 'CALIBRA  ·  CONTROL ESTADÍSTICO DE CALIDAD  ·  Datos del archivo analizado',
         .5, 7.23, 10.85, .18, 8, MUTED)
    if page is not None:
        text(slide, str(page), 12.3, 7.22, .38, .19, 9, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def add_row_table(prs, category, title, header, rows, widths, notes_text,
                  context='', rows_per_slide=12, row_colors=None):
    assert abs(sum(widths)-12.5)<.05
    results=[]
    for page in range(0, max(len(rows),1), rows_per_slide):
        subset=rows[page:page+rows_per_slide]
        slide=titled(prs,category,title + (f' ({page//rows_per_slide +1})' if page else ''),context,len(prs.slides)+1)
        x0=.42
        shape(slide,x0,1.68,12.5,.51,NAVY,radius=False)
        xpos=x0
        for h,w in zip(header,widths):
            text(slide,h,xpos+.09,1.82,w-.15,.22,10,WHITE,True,PP_ALIGN.CENTER)
            xpos+=w
        line_h=min(.54, 4.63/max(len(subset),1))
        for idx, row in enumerate(subset):
            y=2.2+idx*line_h
            shape(slide,x0,y,12.5,line_h, PALE if idx%2 else WHITE, radius=False)
            xpos=x0
            for j,(cell,w) in enumerate(zip(row,widths)):
                color=INK
                if row_colors and (idx+page,j) in row_colors:
                    color=row_colors[(idx+page,j)]
                text(slide,str(cell),xpos+.09,y+line_h*.28,w-.17,line_h*.48,
                     9.7 if len(subset)>10 else 11,color,bool(row_colors and (idx+page,j) in row_colors),
                     PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER)
                xpos+=w
        if not subset:
            text(slide,'No existen filas evaluables para esta sección.', .7,2.5,11.5,.4,13,MUTED)
        notes(slide,notes_text)
        results.append(slide)
    return results


def rows_data(records):
    return [z['result'] for z in records]


def colors_for(r):
    return operational_status(r)[0]


def simple_title(prs, title, records):
    slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid()
    slide.background.fill.fore_color.rgb=rgb(WHITE)
    shape(slide,.0,.0,.13,7.5,NAVY,radius=False)
    text(slide,'CONTROL ESTADÍSTICO DE CALIDAD',.68,1.24,11.8,.37,16,BLUE,True)
    text(slide,title[:75],.66,2.05,12.0,1.22,35,NAVY,True)
    shape(slide,.71,3.54,3.3,.055,BLUE,radius=False)
    n=sum(z['result']['summary']['total'] for z in records)
    fn=list(dict.fromkeys(z['filename'] for z in records))
    text(slide,(fn[0][:94] if len(fn)==1 else f'{len(fn)} archivos analizados'),
         .7,4.04,11.7,.48,18,INK)
    text(slide,f'{len(records)} parámetros · {n:,} mediciones válidas',.7,4.77,11.5,.45,16,INK)
    periods=sorted(set(period_label(z['result']) for z in records if period_label(z['result'])))
    text(slide,(periods[0] if len(periods)==1 else 'Cortes y períodos identificados en los registros')
         if periods else 'Período completo del documento analizado', .7,5.57,11.5,.38,12,MUTED)
    text(slide,'Elaborado automáticamente por Calibra · '+datetime.now().strftime('%d/%m/%Y'),.7,6.75,11.5,.3,11,MUTED)
    notes(slide,[
        'APERTURA. Esta presentación resume únicamente las mediciones válidas archivadas en Calibra. '
        'Los archivos Excel de origen no han sido modificados.',
        f'Se incluyeron {len(records)} parámetros o análisis y {n} mediciones válidas en total. '
        'Si un ensayo aparece en más de un análisis, sus recuentos pueden repetirse; comprobar los nombres y períodos.',
        'El objetivo es distinguir cumplimiento de las especificaciones contractuales y estabilidad estadística del proceso. '
        'Las decisiones de aceptación se deben verificar contra el contrato vigente.'
    ])


def percent(n,total):
    return f'{100*n/total:.1f} %' if total else '—'


def semaforo(prs,records):
    count={k:0 for k in STATUS}
    for z in records:count[colors_for(z['result'])]+=1
    slide=titled(prs,'RESUMEN EJECUTIVO','Resumen del control de calidad',page=len(prs.slides)+1)
    total=sum(z['result']['summary']['total'] for z in records)
    outside=sum(z['result']['summary']['above']+z['result']['summary']['below'] for z in records if z['result']['spec_defined'])
    omitted=sum(z['result']['summary']['omitted'] for z in records)
    values=[('REGISTROS VÁLIDOS',str(total)),('PARÁMETROS',str(len(records))),('FUERA DE ESPEC.',str(outside)),('LECTURAS OMITIDAS',str(omitted))]
    for i,(name,n) in enumerate(values):
        x=.57+i*3.13
        shape(slide,x,1.75,2.91,1.08,PALE,radius=False)
        text(slide,n,x+.08,1.92,2.75,.40,25,NAVY,True,PP_ALIGN.CENTER)
        text(slide,name,x+.1,2.51,2.7,.22,9.3,MUTED,True,PP_ALIGN.CENTER)
    text(slide,'SEMÁFORO DE PARÁMETROS',.6,3.31,11.4,.3,15,NAVY,True)
    for i,(name,color) in enumerate(STATUS.items()):
        x=.55+i*3.15
        shape(slide,x,3.96,2.94,.79,'#EDF2F6',color)
        text(slide,name,x+.12,4.09,2.7,.25,11,color,True,PP_ALIGN.CENTER)
        text(slide,str(count[name]),x+.12,4.39,2.7,.30,20,color,True,PP_ALIGN.CENTER)
    maybe=[z['result']['settings']['metric'] for z in records if not z['result']['spec_defined']]
    text(slide,('Sin límites confirmados: '+', '.join(maybe[:5])) if maybe else 'Todas las variables presentan especificaciones configuradas.',
         .66,5.35,11.95,.55,12,MUTED)
    if outside:
        text(slide,f'ATENCIÓN: {outside} resultados fuera de especificación en los parámetros configurados.',
             .66,6.25,11.9,.35,13,RED,True)
    notes(slide,[
        f'PANORAMA. Se procesaron {total} mediciones en {len(records)} análisis, con {outside} valores fuera de los límites de especificación configurados y {omitted} lecturas sin un valor numérico usable.',
        'El semáforo mostrado es el criterio operativo de Calibra: ROJO significa que hay al menos una lectura fuera de especificación; AMARILLO indica alertas de control u omisiones sin excedencia; VERDE significa que no se detectaron estas incidencias; NO EVALUABLE corresponde a un análisis sin límites configurados.',
        'No interprete estos colores como la validación de un intervalo de tolerancia estadístico 90/90. Son reglas operativas diferentes; el contrato y sus especificaciones tienen prioridad.',
        'Las sumas se refieren a análisis y pueden contabilizar la misma muestra más de una vez cuando se utiliza un compilado de varios archivos o mediciones.'
    ])


def compliance(prs,records):
    rows=[];col={}
    for i,z in enumerate(records):
        r=z['result'];s=r['summary'];o=s['below']+s['above'];spec=r['spec_defined']
        status='Sin límite' if not spec else ('Cumple' if o==0 and s['omitted']==0 else ('Revisar' if o==0 else 'No cumple'))
        rows.append((r['settings']['metric'][:29],limitation(r) if spec else 'Sin límite confirmado',
                     str(s['total']),str(o) if spec else '—',percent(s['within'],s['total']) if spec else '—',status))
        col[(i,5)]=RED if status=='No cumple' else (GREEN if status=='Cumple' else MUTED)
    return add_row_table(prs,'CUMPLIMIENTO','Resultados fuera de especificación',
        ['Parámetro','Especificación','Muestras','Fuera de rango','Cumplimiento','Estado'],
        rows,[3.05,2.35,1.5,1.73,2.0,1.87],
        'CUMPLIMIENTO. Esta tabla coteja cada medición válida contra el mínimo y/o máximo que el usuario configuró para ese análisis. '
        'La cantidad fuera de rango es la suma de resultados por debajo del mínimo y por encima del máximo. '
        'Si no hay límites confirmados se informa «Sin límite», sin aprobar ni rechazar el ensayo. '
        'Si hay lecturas omitidas, el resultado requiere revisión aunque las observadas estén dentro del rango. '
        'Estos resultados no prueban estabilidad estadística; las cartas de control se interpretan por separado.',
        rows_per_slide=10,row_colors=col)


def methods(prs):
    slide=titled(prs,'METODOLOGÍA','Metodología estadística',page=len(prs.slides)+1)
    bullets=[
        'Carta X–mR: σ dentro del proceso = rango móvil promedio / 1.128; límites de individuales = media ± 3σ.',
        'Carta de rangos móviles: el límite superior es 3.267 × MR promedio. Los límites de control no son límites de especificación.',
        'Intervalo bilateral al 90 % de confianza de la media mediante t de Student, cuando hay datos suficientes.',
        'Cpk: usa σ dentro del proceso; Ppk: usa la desviación estándar global. La no normalidad reduce su interpretación.',
        'Señales exploratorias: observaciones fuera de ±3σ, rachas ≥8 del mismo lado y secuencias de 6 ascendentes o descendentes.',
        'Semáforo operativo: rojo con excedencias; amarillo con otras alertas; verde sin alertas; no evaluable sin límites.',
        'La tolerancia estadística 90/90 NO se calcula en esta versión: no se presenta como aprobada o rechazada.'
    ]
    for i,line in enumerate(bullets):
        y=1.63+.69*i
        shape(slide,.57,y,12.04,.58,PALE,radius=False)
        text(slide,'• '+line,.75,y+.09,11.65,.40,12.1,INK)
    notes(slide,[
        'METODOLOGÍA. El gráfico de individuales controla la variación de observaciones consecutivas, asumiendo que el orden representa la secuencia real del proceso. '
        'La dispersión dentro del proceso se estima con el rango móvil de 2 observaciones y la constante d2=1.128. '
        'Para el gráfico MR, el factor del límite superior es D4=3.267.',
        'La especificación viene del contrato o de los límites suministrados por el usuario; es distinta de los límites estadísticos de control estimados a partir de las muestras. '
        'Puede existir cumplimiento de especificación y al mismo tiempo inestabilidad estadística.',
        'Los gráficos de distribución y Q–Q son diagnósticos. Cpk y Ppk no son criterios únicos de aceptación cuando hay distribución no normal o causas especiales. '
        'Calibra no reproduce la prueba de tolerancia 90/90 utilizada por ciertos informes de referencia; no inventa esos resultados.',
        'Se recomienda evaluar la estabilidad por subgrupos homogéneos y revisar los cambios de laboratorio, materiales, dosificación o tiempo.'
    ])


def normalized_lab(label):
    x=unicodedata.normalize('NFKD',str(label or '').casefold())
    x=''.join(c for c in x if not unicodedata.combining(c))
    if 'contratista' in x: return 'Contratista'
    if 'supervision' in x or 'supervisor' in x: return 'Supervisión'
    return None


def lab_stats(r):
    groups={'Contratista':[], 'Supervisión':[]}
    for p in r['points']:
        label=normalized_lab(p.get('laboratory'))
        if label:groups[label].append(float(p['value']))
    return {k:(mean(v),len(v)) if v else None for k,v in groups.items()}


def tendency(r):
    x=[float(p['value']) for p in r['points']]
    avg=float(r['summary']['mean'])
    streak=0; best=0; oldside=0
    for v in x:
        side=1 if v>avg else -1 if v<avg else 0
        streak=streak+1 if side and side==oldside else 1 if side else 0
        best=max(best,streak);oldside=side
    trend=False
    for i in range(len(x)-5):
        seq=x[i:i+6]
        if all(seq[k]<seq[k+1] for k in range(5)) or all(seq[k]>seq[k+1] for k in range(5)):
            trend=True;break
    observed_issues=r['summary']['spc_violations']
    if len(x)<8:
        stable='Indeterminado'
    else:
        stable='No' if observed_issues or r['summary']['mr_violations'] or best>=8 or trend else 'Sin señales'
    return best,trend,observed_issues,stable


def detail_notes(r):
    summary=r['summary'];metric=r['settings']['metric'];a=advanced(r);stat=colors_for(r)
    outside=summary['above']+summary['below'];within=summary['within'];n=summary['total']
    run,trend,out_ctrl,st=tendency(r)
    lab=lab_stats(r)
    pieces=[
        f'PARÁMETRO: {metric}. En el análisis se incluyeron {n} mediciones válidas. Media = {fmt(summary["mean"],3)}; '
        f'mínimo = {fmt(summary["minimum"],3)} y máximo = {fmt(summary["maximum"],3)}. '
        f'Especificación configurada: {limitation(r) if r["spec_defined"] else "SIN LÍMITES CONFIRMADOS"}.',
        (f'CUMPLIMIENTO: {within} de {n} mediciones están dentro de los límites definidos ({fmt(summary["percent"],1)} %). '
         f'{summary["below"]} valores son inferiores al mínimo y {summary["above"]} exceden el máximo.'
         if r['spec_defined'] else
         'CUMPLIMIENTO: no puede concluirse aprobación o rechazo al no existir límites de especificación documentados.'),
        f'CONTROL ESTADÍSTICO: {out_ctrl} lecturas fuera de límites I; {summary["mr_violations"]} saltos fuera del límite MR; '
        f'racha máxima del mismo lado de la media = {run}, tendencia de 6 lecturas = {"sí" if trend else "no"}. '
        f'Estabilidad operativa: {st}. Una lectura conforme puede igualmente señalar una causa especial.',
        f'DIAGNÓSTICO: Cpk = {fmt(a["cpk"],3)}; Ppk = {fmt(a["ppk"],3)}; '
        f'Shapiro p = {fmt(a["shapiro"],4)}; IC 90 % para la media = '
        f'[{fmt(a["ci"][0],3)}, {fmt(a["ci"][1],3)}]. '
        'Las estimaciones de capacidad bajo no normalidad o inestabilidad deben interpretarse con cautela.',
        f'TRAZABILIDAD: el resultado conserva identificadores de fila original y período seleccionado. '
        f'Hay {summary["omitted"]} filas sin una medición numérica usable. '
        'Se recomienda revisar fechas, registros atípicos, dosificación y cambios de proceso antes de concluir causas.',
    ]
    if all(lab.values()):
        cc,ss=lab['Contratista'],lab['Supervisión']
        pieces.append(f'LABORATORIOS: Contratista media {cc[0]:.3f} (n={cc[1]}), '
                      f'Supervisión media {ss[0]:.3f} (n={ss[1]}). Diferencia no pareada = {cc[0]-ss[0]:+.3f}; '
                      'una diferencia de medias no demuestra sesgo estadístico por sí sola.')
    if n<20:
        pieces.append('ADVERTENCIA DE TAMAÑO MUESTRAL: menos de 20 observaciones. '
                      'Los límites de control estimados pueden ser poco representativos y deben tratarse como exploratorios.')
    return pieces



# Semáforo y comentario visible, inspirado en la lámina enviada por el usuario.
# Esta nota NO reemplaza las notas del presentador: se lee directamente en la diapositiva.
def visible_result_note(prs, record, idx):
    r=record['result']; s=r['summary']; metric=r['settings']['metric']
    code=colors_for(r); accent=STATUS[code]
    pastel={'VERDE':'#E7F5EC','AMARILLO':'#FFF2D6',
            'ROJO':'#FCEAEA','NO EVALUABLE':'#EDF1F6'}[code]
    n=s['total']; out=s['above']+s['below']; om=s['omitted']
    run,trend,out_control,stability=tendency(r)
    adv=advanced(r)

    if not r['spec_defined']:
        tag='NO EVALUABLE'
        main='No se han confirmado los límites de especificación para este parámetro; no se puede determinar si cumple.'
        action='Confirmar los límites contractuales aplicables antes de emitir un dictamen de conformidad.'
    elif out:
        tag='ALERTA'
        main=(f'{out} de {n} resultados están fuera de la especificación '
              f'({fmt(s["percent"],1)} % de cumplimiento).')
        detail=[]
        if s['below']: detail.append(f'{s["below"]} por debajo del mínimo')
        if s['above']: detail.append(f'{s["above"]} por encima del máximo')
        action='Investigar las mediciones fuera de rango: '+', '.join(detail)+'.'
    elif om:
        tag='REVISIÓN PENDIENTE'
        main=f'Las {n} mediciones numéricas están dentro de la especificación (100 % de las válidas).'
        action=f'Se omitieron {om} registro(s) no numéricos; revisarlos antes de confirmar la integridad del ensayo.'
    else:
        tag='CUMPLIMIENTO'
        main=f'Todos los resultados ({n}) están dentro de la especificación (100 %).'
        action='No se detectaron lecturas fuera de los límites de especificación configurados.'

    # Las señales de estabilidad se describen sin confundirlas con excedencias contractuales.
    signals=[f'{out_control} punto(s) fuera de control',f'{s["mr_violations"]} salto(s) MR',
             f'racha máxima {run}']
    if trend: signals.append('tendencia de 6 mediciones')
    if stability=='Sin señales':
        stability_line='No se detectaron señales de inestabilidad con las reglas implementadas.'
    elif stability=='Indeterminado':
        stability_line='Estabilidad indeterminada: se requieren más datos para evaluar rachas de forma confiable.'
    else:
        stability_line='Señales de inestabilidad: '+', '.join(signals)+'.'
    if n<20: stability_line+=' La muestra es reducida; interpretar las cartas con cautela.'
    diagnostic=[]
    if adv['shapiro'] is not None and adv['shapiro']<.05:
        diagnostic.append('distribución no normal (Shapiro p < 0.05)')
    if stability=='No': diagnostic.append('proceso con señales de inestabilidad')
    if not r['spec_defined']:
        caution='Cpk no calculable sin especificaciones; el control estadístico sigue siendo orientativo.'
    elif diagnostic:
        caution='Cpk orientativo: '+', '.join(diagnostic)+'.'
    else:
        caution='Cpk referencial: corroborar la validez de los límites y del muestreo antes de decidir aceptación.'

    slide=titled(prs,f'ANÁLISIS {idx:02d} — NOTAS VISIBLES',
                 f'{metric} · observaciones de calidad',
                 f'Hoja: {r["settings"]["sheet"]} · Fuente: {record["filename"][:65]}  '
                 + period_label(r),len(prs.slides)+1)

    # Emula el estado, las tres cifras y el recuadro de la captura de referencia.
    shape(slide,.76,1.66,2.03,.52,pastel,accent)
    text(slide,code,.88,1.79,1.79,.22,11.3,accent,True,PP_ALIGN.CENTER)
    text(slide,'Espec. '+(limitation(r) if r['spec_defined'] else 'sin límites confirmados'),
         3.01,1.80,9.22,.29,12.3,MUTED)
    for x,label,value,metric_color in [
        (1.15,'muestras (n)',str(n),NAVY),
        (4.95,'media',fmt(s['mean'],3),NAVY),
        (8.89,'Cpk',fmt(adv['cpk'],2),accent)]:
        text(slide,value,x,2.38,2.68,.46,23.5,metric_color,True,PP_ALIGN.CENTER)
        text(slide,label,x,2.94,2.68,.25,10.6,MUTED,False,PP_ALIGN.CENTER)

    shape(slide,.75,3.51,11.83,2.71,pastel,accent)
    text(slide,tag,1.04,3.82,11.04,.34,15,accent,True)
    text(slide,main,1.04,4.32,10.96,.55,14.6,INK)
    text(slide,stability_line,1.04,4.98,10.96,.62,12.6,INK)
    text(slide,action,1.04,5.74,10.96,.40,10.8,INK)
    text(slide,caution,.89,6.49,11.7,.35,9.7,MUTED)

    notes(slide, [
        f'RECUADRO VISIBLE — {tag}. {main}',
        stability_line, action, caution,
        *detail_notes(r),
    ])
    return slide


def metric_slides(prs,record,idx):
    r=record['result'];s=r['summary'];a=advanced(r);metric=r['settings']['metric']
    status=colors_for(r);col=STATUS[status]
    slide=titled(prs,f'ANÁLISIS {idx:02d} — ESTADÍSTICO',f'{metric} · control y cumplimiento',
                 f'Hoja: {r["settings"]["sheet"]} · Fuente: {record["filename"][:65]}  '
                 + period_label(r),len(prs.slides)+1)
    xs=[.58,3.16,5.66,8.13,10.54]
    vals=[('ESTADO',status),('ESPEC.',limitation(r)),('MUESTRAS',str(s['total'])),
          ('MEDIA',fmt(s['mean'],3)),('CPK',fmt(a['cpk'],2))]
    widths=[2.39,2.28,2.24,2.15,2.15]
    for i,((label,value),x,w) in enumerate(zip(vals,xs,widths)):
        shape(slide,x,1.52,w,.63,PALE,radius=False)
        text(slide,label,x+.08,1.58,w-.13,.17,8.5,MUTED,True)
        text(slide,value[:23],x+.08,1.77,w-.14,.27,13.3,col if i==0 else NAVY,True)
    # Chart proportions modeled on reference: individual chart stacked over distribution
    text(slide,'Gráfico de control (individuales, X–mR)',.64,2.27,11.94,.25,13.8,NAVY,True)
    im=make_plot(r,'I',figsize=(11.9,2.75))
    slide.shapes.add_picture(im,Inches(.66),Inches(2.55),width=Inches(12.0),height=Inches(1.65))
    text(slide,'Histograma · distribución de mediciones',.64,4.36,11.94,.25,13.8,NAVY,True)
    im=make_plot(r,'H',figsize=(11.9,2.75))
    slide.shapes.add_picture(im,Inches(.66),Inches(4.68),width=Inches(12.0),height=Inches(1.78))
    out=s['below']+s['above']
    warning=(f'ALERTA: {out} de {s["total"]} resultados fuera de especificación.' if out
             else ('Cumplimiento no evaluable: límites sin confirmar.' if not r['spec_defined']
                   else f'CUMPLIMIENTO: {s["total"]} lecturas dentro de especificación.'))
    text(slide,warning,.7,6.73,11.9,.30,10.3,RED if out else (MUTED if not r['spec_defined'] else GREEN),True)
    notes(slide, detail_notes(r))
    # Diapositiva legible con el recuadro de CUMPLIMIENTO / ALERTA pedido.
    visible_result_note(prs, record, idx)
    # Preserve the previous detailed mR and Q-Q charts, in a separate complement.
    slide2=titled(prs,'ANÁLISIS COMPLEMENTARIO',f'{metric} · estabilidad y diagnóstico',
                  'Rangos móviles (MR) y prueba visual de normalidad (Q–Q).',len(prs.slides)+1)
    for x,y,w,h,k,heading in [(0.59,1.83,12.0,2.2,'MR','Rangos móviles MR y límite superior'),
                              (0.59,4.41,12.0,2.2,'QQ','Gráfico de probabilidad normal Q–Q')]:
        text(slide2,heading,x,y-.28,w,.26,14,NAVY,True)
        im=make_plot(r,k,figsize=(11.9,2.75))
        slide2.shapes.add_picture(im,Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    notes(slide2,['DIAGNÓSTICO COMPLEMENTARIO.',
        *detail_notes(r),
        'En la carta MR cada punto representa el cambio absoluto entre dos mediciones consecutivas. '
        'En el gráfico Q–Q las desviaciones respecto a la línea recta sugieren que la distribución de mediciones puede apartarse de la normal. '
        'Este diagnóstico no sustituye el estudio por fecha, estación o laboratorio.'])



def incidents(prs, record, idx):
    r=record['result']
    items=[]
    for p in r['points']:
        tags=[]
        if r['spec_defined'] and not p['in_specs']: tags.append('Fuera de especificación')
        if not p['in_control']: tags.append('Señal fuera de control I')
        if p.get('mr') is not None and p['mr']>r['limits']['mr_ucl']: tags.append('Salto MR')
        if tags:
            items.append((str(p['row']),str(p['label'])[:43],fmt(p['value'],4),'; '.join(tags)))
    for o in r.get('omitted',[]):
        items.append((str(o['row']),str(o['label'])[:43],str(o.get('value',''))[:16], 'Lectura omitida'))
    if not items:
        return
    colors={(i,3):RED for i,it in enumerate(items) if 'Fuera de especificación' in it[3]}
    add_row_table(prs, 'INCIDENCIAS', f'Incidencias detectadas · {r["settings"]["metric"][:46]}',
        ['Fila Excel','Etiqueta o fecha','Medición','Motivo / tipo de incidencia'],
        items,[1.38,3.68,1.67,5.77],
        f'INCIDENCIAS DEL ANÁLISIS {idx}. Se listan {len(items)} filas con alguna señal estadística, fuera de especificación o sin valor válido. '
        'La columna de fila corresponde a la ubicación registrada en el Excel original. '
        'Un mismo registro puede activar varias reglas y aparecer agrupado en una sola fila de la lista. '
        'Un punto fuera de control no siempre está fuera de la especificación contractual; diagnosticar ambas señales por separado.',
        context='Trazabilidad de filas del Excel original; no se modificaron los datos de origen.',
        rows_per_slide=13, row_colors=colors)


def laboratories(prs,records):
    rows=[]
    for z in records:
        r=z['result'];stats=lab_stats(r);c=stats['Contratista'];sup=stats['Supervisión']
        if not c and not sup:continue
        rows.append((r['settings']['metric'][:31],
                     f'{c[0]:.3f} (n={c[1]})' if c else 'Sin registros',
                     f'{sup[0]:.3f} (n={sup[1]})' if sup else 'Sin registros',
                     f'{c[0]-sup[0]:+.3f}' if c and sup else '—'))
    add_row_table(prs,'LABORATORIOS','Contratista frente a Supervisión',
        ['Parámetro','Contratista (media · n)','Supervisión (media · n)','Diferencia'],
        rows,[3.31,3.55,3.55,2.09],
        'LABORATORIOS. La diferencia se calcula como media del Contratista menos media de Supervisión, con base únicamente en la columna «Laboratorio» de cada fila del Excel. '
        'Esta es una comparación descriptiva no pareada: cada laboratorio puede haber ensayado muestras y fechas distintas. '
        'Una diferencia persistente motiva revisión de trazabilidad, método, calibración y coincidencia temporal, pero no demuestra sesgo por sí sola. '
        'La ausencia de registros de un laboratorio se informa explícitamente; no se estiman promedios ficticios.',
        context='Diferencia = media Contratista − media Supervisión; no implica por sí sola sesgo estadístico.',rows_per_slide=10)


def stability(prs,records):
    rows=[]; colors={}
    for i,z in enumerate(records):
        r=z['result'];streak,trend,out,st=tendency(r)
        rows.append((r['settings']['metric'][:30],str(streak),'Sí' if trend else 'No',str(out),st))
        colors[(i,4)]=RED if st=='No' else (GREEN if st=='Sin señales' else MUTED)
    add_row_table(prs,'ESTABILIDAD','Tendencias y cambios de proceso',
        ['Parámetro','Racha máxima','Tendencia (6)','Fuera de control','Proceso estable'],
        rows,[3.25,2.32,2.23,2.52,2.18],
        'ESTABILIDAD. Racha máxima significa el mayor número de observaciones consecutivas del mismo lado de la media (los empates interrumpen la racha). '
        'Tendencia se marca cuando existen 6 mediciones estrictamente ascendentes o descendentes. '
        'Fuera de control muestra lecturas por encima del LCS o por debajo del LCI del gráfico de individuales. '
        'La columna «Proceso estable» es exploratoria: «No» señala al menos una regla activada en individuales, rangos móviles, rachas o tendencia; '
        '«Sin señales» no demuestra estabilidad futura, e «Indeterminado» corresponde a una muestra inferior a ocho.',
        context='Una medición dentro de especificación puede coexistir con inestabilidad estadística.',rows_per_slide=11,row_colors=colors)


def dashboard(prs,records):
    for page in range(0,len(records),8):
        subset=records[page:page+8]
        slide=titled(prs,'DASHBOARD','Dashboard de calidad' + (f' · {page//8+1}' if page else ''),page=len(prs.slides)+1)
        for j,z in enumerate(subset):
            r=z['result'];s=r['summary'];st=colors_for(r);color=STATUS[st]
            rr,cc=divmod(j,4)
            x=.52+cc*3.18;y=1.72+rr*2.41
            shape(slide,x,y,3.02,2.15,PALE,radius=False)
            shape(slide,x,y,.065,2.15,color,radius=False)
            text(slide,r['settings']['metric'][:27],x+.18,y+.19,2.75,.32,12,NAVY,True)
            text(slide,st,x+.18,y+.66,2.67,.27,11,color,True)
            text(slide,'Cumplimiento '+(fmt(s['percent'],1)+' %' if r['spec_defined'] else 'no evaluable'),x+.18,y+1.11,2.75,.25,10,INK)
            text(slide,f'Cpk {fmt(advanced(r)["cpk"],2)} · n={s["total"]}',x+.18,y+1.58,2.68,.26,10,MUTED)
        notes(slide,['DASHBOARD DE CALIDAD.',
            f'Se muestran los parámetros {page+1} a {page+len(subset)} de {len(records)}. '
            'Cada tarjeta expresa el semáforo operativo, el porcentaje empírico de cumplimiento y el Cpk cuando es calculable.',
            'No compare Cpk entre procesos sin verificar primero normalidad, estabilidad y homogeneidad de datos. '
            'El porcentaje observado no equivale a probabilidad de cumplimiento futuro ni a una garantía contractual.'])


def recommendations(prs,records):
    out=[z for z in records if z['result']['spec_defined'] and z['result']['summary']['above']+z['result']['summary']['below']]
    unclear=[z for z in records if not z['result']['spec_defined']]
    unstable=[z for z in records if tendency(z['result'])[3]=='No']
    omitted=sum(z['result']['summary']['omitted'] for z in records)
    todo=[]
    if out:todo.append('Investigar resultados fuera de especificación en: '+', '.join(z['result']['settings']['metric'] for z in out)[:160]+'.')
    if unclear:todo.append('Confirmar documentalmente especificaciones de: '+', '.join(z['result']['settings']['metric'] for z in unclear)[:160]+'.')
    if unstable:todo.append('Analizar causas especiales, fechas y laboratorios en: '+', '.join(z['result']['settings']['metric'] for z in unstable)[:150]+'.')
    if omitted:todo.append(f'Revisar {omitted} lectura(s) omitida(s), celdas vacías o registros no numéricos antes de aprobar datos.')
    todo.extend(['Verificar la trazabilidad por estación, turno, dosificación y laboratorio cuando existan cambios en el proceso.',
                 'No utilizar Cpk, un histograma o el porcentaje observado como únicos criterios de aceptación.'])
    slide=titled(prs,'RECOMENDACIONES','Acciones correctivas recomendadas',page=len(prs.slides)+1)
    for i,line in enumerate(todo[:7]):
        y=1.66+i*.69
        shape(slide,.57,y,12.06,.58,PALE,radius=False)
        text(slide,'• '+line,.75,y+.105,11.64,.36,11.4,INK)
    notes(slide,['ACCIONES PROPUESTAS. Las recomendaciones se originan en las incidencias observadas; '
        'no son decisiones finales de aceptación de obra.',
        'Frente a cualquier excedencia, verificar la especificación aplicable, la calibración de equipos, '
        'la validez del dato y la localización de la muestra antes de decidir una acción correctiva.',
        'Cuando existan diferencias entre laboratorios, priorizar comparaciones de muestras emparejadas y el contexto operativo. '
        'Documentar investigación y decisiones bajo el procedimiento de control de calidad vigente.'])


def conclusions(prs,records):
    n=sum(z['result']['summary']['total'] for z in records)
    out=[z for z in records if z['result']['spec_defined'] and (z['result']['summary']['above']+z['result']['summary']['below'])]
    o=sum(z['result']['summary']['above']+z['result']['summary']['below'] for z in out)
    unknown=[z for z in records if not z['result']['spec_defined']]
    lines=[
        f'El informe se sustenta en {n:,} mediciones válidas de {len(records)} parámetros/analíticas archivadas.',
        f'Se identificaron {o} observaciones fuera de especificación' + (': '+', '.join(z['result']['settings']['metric'] for z in out)[:148] if out else '.') ,
        ('Sin límite confirmado: '+', '.join(z['result']['settings']['metric'] for z in unknown)[:156]+'.') if unknown else 'Todos los ensayos cuentan con límites configurados.',
        'Las cartas distinguen cumplimiento de estabilidad; ambas evaluaciones deben revisarse conjuntamente.',
        'El Excel de origen permanece sin alteraciones y las incidencias son trazables a sus filas.',
    ]
    slide=titled(prs,'CIERRE','Conclusiones',page=len(prs.slides)+1)
    for i,line in enumerate(lines):
        y=1.72+i*.94
        shape(slide,.58,y,12.05,.79,PALE,radius=False)
        text(slide,'• '+line,.79,y+.17,11.62,.47,12.4,INK)
    notes(slide,['CIERRE. Los resultados corresponden al contenido efectivamente procesado y a los límites introducidos o configurados en Calibra.',
        f'Fueron analizados {n} registros válidos y se encontraron {o} lecturas fuera de especificación. '
        'Evitar interpretar estos resultados como aceptación global de todas las características del proyecto si el Excel contiene ensayos no incluidos.',
        'Todo análisis sin límite confirmado queda expresamente marcado como no evaluable; no se inventan umbrales. '
        'El responsable técnico debe validar límites, períodos, ensayos y representatividad antes de usar el informe en decisiones contractuales.',
        'Las notas de cada diapositiva aportan contexto para presentar y discutir hallazgos en reuniones de seguimiento.'])


def build_presentation(records,title='Informe de control estadístico',groups=None):
    records=[normalize_record(z) for z in records]
    if not records:raise ValueError('No hay resultados archivados para generar un PowerPoint.')
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    prs.core_properties.title=title;prs.core_properties.subject='Calibra · informe de calidad con notas de presentación';prs.core_properties.author='Calibra'
    simple_title(prs,title,records)
    semaforo(prs,records)
    compliance(prs,records)
    methods(prs)
    if groups:
        for gi,(info,group_records) in enumerate(groups,1):
            subset=[normalize_record(z) for z in group_records]
            if not subset: continue
            slide=titled(prs,'ARCHIVO EN BIBLIOTECA',f'{gi:02d} · {info.get("filename", "Archivo")[:74]}',
                         f'{len(subset)} parámetros · informe individual dentro del compilado',len(prs.slides)+1)
            n=sum(z['result']['summary']['total'] for z in subset)
            text(slide,f'{n} mediciones válidas',.85,2.48,10.7,.65,26,NAVY,True)
            text(slide,'La sección siguiente contiene solo los ensayos de este archivo.',.87,3.58,11.1,.44,14,MUTED)
            notes(slide,[f'INICIO DEL ARCHIVO {gi}. Los datos corresponden a {info.get("filename", "archivo sin nombre")}.',
                              'Al compilar varios archivos puede existir solapamiento de muestras; distinguir los resultados de cada corte para evitar dobles recuentos.'])
            semaforo(prs,subset);compliance(prs,subset)
            for k,z in enumerate(subset,1):
                metric_slides(prs,z,k)
                incidents(prs,z,k)
            laboratories(prs,subset);stability(prs,subset);dashboard(prs,subset)
    else:
        for i,z in enumerate(records,1):
            metric_slides(prs,z,i)
            incidents(prs,z,i)
        laboratories(prs,records);stability(prs,records);dashboard(prs,records)
    recommendations(prs,records)
    conclusions(prs,records)
    # Diagnostics should be bounded: native notes exist on EVERY slide.
    assert all(slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip() for slide in prs.slides)
    memory=io.BytesIO();prs.save(memory);memory.seek(0)
    return memory
