"""Generación reproducible de presentaciones para análisis guardados en Calibra.

El archivo fuente no se abre ni se modifica: se utilizan exclusivamente los
resultados JSON calculados por la aplicación y almacenados en la base de datos.
"""
from __future__ import annotations

import io
import math
from datetime import datetime
from statistics import stdev

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.ticker import MaxNLocator
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from scipy import stats

NAVY = '#192D40'
INK = '#263F54'
MUTED = '#64798B'
TEAL = '#0C9E98'
GREEN = '#28754C'
RED = '#B94048'
AMBER = '#B6811C'
BLUE = '#4E789B'
PALE = '#F4F7F9'
LINE = '#DFE8EC'
WHITE = '#FFFFFF'
ORANGE = '#C8920D'


def rgb(value):
    return RGBColor.from_string(value.lstrip('#'))


def text(slide, value, x, y, w, h, size=14, color=INK, bold=False, align=None,
         font='Aptos', margin=0.03, lines=1, valign=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    if valign is not None:
        frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.text = str(value)
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bool(bold)
    p.font.color.rgb = rgb(color)
    if align is not None:
        p.alignment = align
    return box


def shape(slide, x, y, w, h, fill=WHITE, line=None, radius=True):
    typ = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    obj = slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(fill)
    obj.line.fill.background() if not line else None
    if line:
        obj.line.color.rgb = rgb(line)
        obj.line.width = Pt(.7)
    if radius:
        try: obj.adjustments[0] = 0.1
        except Exception: pass
    return obj


def base(prs, category, title, subtitle='', page=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    shape(slide, 0, 0, 13.333, .11, TEAL, radius=False)
    text(slide, category.upper(), .51, .27, 10, .24, 9, TEAL, True)
    text(slide, title, .51, .6, 12.2, .54, 25, NAVY, True)
    if subtitle:
        text(slide, subtitle, .53, 1.14, 11.95, .35, 11.2, MUTED)
    shape(slide, .5, 7.17, 12.3, .012, LINE, radius=False)
    text(slide, 'CALIBRA  /  INFORME AUTOMÁTICO DE CONTROL DE CALIDAD', .53, 7.24, 9, .16, 7.8, MUTED)
    if page is not None:
        text(slide, f'{page:02d}', 12.28, 7.22, .46, .2, 9, TEAL, True, PP_ALIGN.RIGHT)
    return slide


def fmt(x, digits=2):
    if x is None or not math.isfinite(float(x)):
        return '—'
    return f'{float(x):,.{digits}f}'.replace(',', ' ') 


def period_label(result):
    period = result.get('settings', {}).get('date_filter')
    if not period:
        return ''
    return f'Período: {period.get("start") or "sin inicio"} a {period.get("end") or "sin fin"}'


def limitation(result):
    s = result['settings']; lo,hi = s.get('lsl'),s.get('usl')
    if lo is None and hi is None: return 'NO DEFINIDA'
    if lo is None: return f'≤ {fmt(hi, 3)}'
    if hi is None: return f'≥ {fmt(lo, 3)}'
    return f'{fmt(lo, 3)} – {fmt(hi, 3)}'


def operational_status(r):
    s = r['summary']
    if not r.get('spec_defined'): return 'NO EVALUABLE', MUTED
    if s['below'] + s['above']: return 'ROJO', RED
    if s['omitted'] or s['spc_violations'] or s['mr_violations']: return 'AMARILLO', AMBER
    return 'VERDE', GREEN


def advanced(r):
    x = np.array([p['value'] for p in r['points']], dtype=float)
    n = x.size
    s = r['settings']
    xbar = float(x.mean())
    mean_sigma = float(np.std(x, ddof=1)) if n > 1 else 0
    sigma_within = float(r['limits']['mr_cl'] / 1.128)
    def capability(sig):
        if not sig or sig <= 0: return None
        candidates = []
        if s['lsl'] is not None: candidates.append((xbar - s['lsl']) / (3 * sig))
        if s['usl'] is not None: candidates.append((s['usl'] - xbar) / (3 * sig))
        return min(candidates) if candidates else None
    shapiro = None
    if 3 <= n <= 5000 and len(set(x)) > 2:
        try: shapiro = float(stats.shapiro(x).pvalue)
        except ValueError: pass
    if n > 1 and mean_sigma > 0:
        t = float(stats.t.ppf(.95, n - 1))
        conf = (xbar - t * mean_sigma/math.sqrt(n), xbar + t * mean_sigma/math.sqrt(n))
    else:
        conf = (xbar, xbar)
    return {'cpk': capability(sigma_within), 'ppk': capability(mean_sigma),
            'shapiro': shapiro, 'ci': conf, 'sigma_within': sigma_within, 'sigma_total': mean_sigma}


def add_kpi(slide, x, y, w, label, value, accent=INK, foot=''):
    shape(slide, x, y, w, .97, PALE, LINE)
    text(slide, label.upper(), x + .14, y+.12, w-.28, .18, 8.7, MUTED, True)
    text(slide, value, x+.13, y+.33, w-.28, .4, 21, accent, True)
    if foot: text(slide, foot, x+.15, y+.77, w-.3, .14, 7.3, MUTED)


def make_plot(r, kind):
    """Render a chart from ALL points (outliers always plotted)."""
    data = r['points']; x = np.asarray([p['value'] for p in data], dtype=float)
    n = x.size; ix = np.arange(1, n + 1)
    s = r['settings']; l = r['limits']
    fig, ax = plt.subplots(figsize=(7.0, 3.35), dpi=154)
    fig.patch.set_facecolor('white'); ax.set_facecolor('#F8FAFC')
    for sp in ax.spines.values(): sp.set_color('#D5DEE5')
    ax.tick_params(colors='#536778', labelsize=8)
    ax.grid(color='#E5EDF2', alpha=.8, lw=.6); ax.set_axisbelow(True)
    if kind == 'I':
        # series with a readable line; for >1500 samples retain extremes per block
        if n <= 1500: subset = ix - 1
        else:
            every = max(1, math.ceil(n / 700)); candidates = {0,n-1}
            for start in range(0,n,every):
                block = x[start:start+every]
                candidates.add(start+int(np.argmin(block)))
                candidates.add(start+int(np.argmax(block)))
            subset = np.asarray(sorted(candidates))
        ax.plot(ix[subset], x[subset], color=BLUE, lw=1.3, zorder=2)
        out = np.array([not p['in_specs'] for p in data], bool)
        bad_ctrl = np.array([not p['in_control'] for p in data], bool)
        ax.scatter(ix[~out & ~bad_ctrl], x[~out & ~bad_ctrl], s=10, c=GREEN, zorder=3)
        ax.scatter(ix[~out & bad_ctrl], x[~out & bad_ctrl], s=19, c=AMBER, zorder=4)
        ax.scatter(ix[out], x[out], s=26, c=RED, zorder=5)
        ax.axhline(l['i_cl'], color=NAVY, lw=1.4, label='Media')
        ax.axhline(l['i_lcl'], color=BLUE, ls='--', lw=1, label='Límites de control')
        ax.axhline(l['i_ucl'], color=BLUE, ls='--', lw=1)
        if s['lsl'] is not None: ax.axhline(s['lsl'], color=RED, lw=1.3, ls=(0,(5,3)), label='Especificación')
        if s['usl'] is not None: ax.axhline(s['usl'], color=RED, lw=1.3, ls=(0,(5,3)), label='Especificación' if s['lsl'] is None else None)
        ax.set_xlabel('Orden de las mediciones', fontsize=8)
        ax.set_ylabel('Valor observado', fontsize=8)
        ax.set_xlim(.5, max(n+.5,3))
    elif kind == 'MR':
        move = np.abs(np.diff(x)); ids=np.arange(2,n+1)
        ax.plot(ids, move, color=BLUE, lw=1.3, marker='o' if len(move)<70 else None, markersize=2)
        ax.axhline(l['mr_cl'],color=NAVY,lw=1.3,label='MR promedio')
        ax.axhline(l['mr_ucl'],color=RED,ls='--',lw=1.2,label='Límite superior MR')
        violations=move>l['mr_ucl']
        ax.scatter(ids[violations],move[violations],c=RED,s=23,zorder=3)
        ax.set_ylim(bottom=0);ax.set_xlim(1,max(n+1,3))
        ax.set_xlabel('Observación actual',fontsize=8);ax.set_ylabel('Rango móvil',fontsize=8)
    elif kind == 'H':
        count = min(35, max(6, round(math.sqrt(n))))
        ax.hist(x,bins=count,density=True,color='#B9CBD5',edgecolor=WHITE,alpha=.93)
        sigma = float(x.std(ddof=1)) if n>1 else 0
        if sigma>0:
            lim_min = min(x.min(), s['lsl'] if s['lsl'] is not None else x.min())
            lim_max = max(x.max(), s['usl'] if s['usl'] is not None else x.max())
            grid=np.linspace(min(lim_min,x.mean()-3.6*sigma),max(lim_max,x.mean()+3.6*sigma),250)
            ax.plot(grid,stats.norm.pdf(grid,x.mean(),sigma),color=NAVY,lw=2,label='Normal ajustada (referencial)')
        if s['lsl'] is not None: ax.axvline(s['lsl'],color=RED,ls='--',lw=1.35)
        if s['usl'] is not None: ax.axvline(s['usl'],color=RED,ls='--',lw=1.35)
        ax.set_ylabel('Densidad',fontsize=8);ax.set_xlabel('Medición',fontsize=8)
    elif kind == 'QQ':
        if n>1 and np.std(x)>0:
            theor, ordered = stats.probplot(x, dist='norm', fit=False)
            ax.scatter(theor,ordered,color=BLUE,s=13,alpha=.85)
            (slope,intercept,rr)=stats.probplot(x,dist='norm',fit=True)[1]
            arr=np.array([min(theor),max(theor)])
            ax.plot(arr,intercept+slope*arr,color=RED,lw=1.3)
        else:
            ax.scatter([0]*n,x,c=BLUE,s=20)
        ax.set_xlabel('Cuantiles normales teóricos',fontsize=8)
        ax.set_ylabel('Mediciones ordenadas',fontsize=8)
    ax.yaxis.set_major_locator(MaxNLocator(6))
    fig.subplots_adjust(left=.095,right=.985,bottom=.16,top=.96)
    out = io.BytesIO();fig.savefig(out,format='png',dpi=154,facecolor='white')
    plt.close(fig);out.seek(0)
    return out


def plot_card(slide, r, kind, x, y, w, h, title_line):
    shape(slide,x,y,w,h,WHITE,LINE)
    text(slide,title_line,x+.16,y+.1,w-.3,.26,13,NAVY,True)
    im=make_plot(r,kind)
    slide.shapes.add_picture(im, Inches(x+.14), Inches(y+.48), width=Inches(w-.28), height=Inches(h-.61))


def normalize_record(row):
    if isinstance(row,dict):
        r = row
        return {'filename':r.get('filename','Archivo sin nombre'),
                'created_at':r.get('created_at',''), 'result':r['result'],
                'id':r.get('id')}
    return {'filename':row.filename, 'created_at':row.created_at.isoformat(), 'result':row.result, 'id':row.id}


def add_title(prs, name, records):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(NAVY)
    shape(slide,0,0,.14,7.5,TEAL,radius=False)
    shape(slide,.65,.72,3.25,.4,'#244559',radius=True)
    text(slide,'CALIBRA  /  CONTROL ESTADÍSTICO',.83,.82,3.0,.19,10,'#B2EFE7',True)
    text(slide,'INFORME DE\nCALIDAD',.84,1.63,11.4,1.65,40,WHITE,True)
    text(slide,name,.86,3.68,11.7,.86,21,'#B9E9E1',True)
    shape(slide,.87,5.1,2.15,.05,TEAL,radius=False)
    text(slide,f'{len(records)} análisis incluidos · {sum(z["result"]["summary"]["total"] for z in records):,} mediciones procesadas',.87,5.48,10,.4,15,WHITE)
    periods = sorted(set(period_label(z['result']) for z in records if period_label(z['result'])))
    if periods:
        text(slide, periods[0] if len(periods) == 1 else 'Períodos distintos por archivo / ensayo',
             .88,6.36,11,.28,13,'#B9E9E1',True)
    text(slide,f'Generado el {datetime.now().strftime("%d/%m/%Y %H:%M")}  ·  Origen: resultados guardados en Calibra',.88,6.9,11,.25,10,'#ACBDCA')


def add_dashboard(prs, records):
    total=len(records)
    n=sum(z['result']['summary']['total'] for z in records)
    out=sum(z['result']['summary']['above']+z['result']['summary']['below'] for z in records)
    pending=sum(not z['result']['spec_defined'] for z in records)
    slide=base(prs,'Resumen ejecutivo','Panorama general de calidad',
               'Comparativo de análisis guardados; puede haber mediciones repetidas si se analizó más de una vez un archivo.')
    width=2.9
    for i,(label,val,col) in enumerate([
        ('Análisis incluidos',str(total),NAVY),('Mediciones válidas',f'{n:,}',TEAL),
        ('Fuera de especificación',f'{out:,}',RED if out else GREEN),('Sin límites definidos',str(pending),AMBER if pending else GREEN)]):
        add_kpi(slide,.53+i*3.1,1.76,width,label,val,col)
    text(slide,'TABLERO DE PARÁMETROS',.56,3.04,10,.23,11,TEAL,True)
    columns=[.63,2.85,6.21,8.05,9.6,11.1]
    heads=['N.º','Medición / hoja','Muestras','Cumplimiento','Fuera','Estado']
    shape(slide,.52,3.38,12.3,.43,NAVY,radius=False)
    for x,label in zip(columns,heads):text(slide,label,x,3.48,(2 if label=='Medición / hoja' else 1.4),.23,9.5,WHITE,True)
    max_rows=11
    for i,z in enumerate(records[:max_rows]):
        r=z['result'];stat,c=operational_status(r);v=r['summary'];y=3.84+i*.266
        if i%2==0:shape(slide,.52,y-.02,12.3,.265,PALE,radius=False)
        vals=[str(i+1),(r['settings']['metric'][:22]),str(v['total']),f'{fmt(v["percent"],1)}%' if r['spec_defined'] else 'N/D',str(v['above']+v['below']) if r['spec_defined'] else 'N/D',stat]
        widths=[.8,3.2,1.5,1.5,1.4,1.6]
        for j,(x,vv) in enumerate(zip(columns,vals)):
            text(slide,vv,x,y+.02,widths[j],.21,9.4,c if j==5 else INK,j==5)
    if len(records)>max_rows:
        text(slide,f'El tablero muestra 11 análisis; los {total} se detallan en las siguientes diapositivas.',.61,6.9,11,.2,9,AMBER)


STATUS_PALETTE = [
    ('VERDE', GREEN, '#E4F3EA'),
    ('AMARILLO', AMBER, '#FFF2D5'),
    ('ROJO', RED, '#FBE8E6'),
    ('NO EVALUABLE', '#61708A', '#E8EDF4'),
]


def add_semaforo(prs, records, heading='Semáforo de parámetros', subtitle=None):
    """Distribución de estados operativos estilo tarjetas del informe de referencia.

    Estos estados son las reglas de Calibra, no una equivalencia con
    tolerancia 90/90 ni certificación del laboratorio.
    """
    records = [normalize_record(z) for z in records]
    counts = {name: 0 for name, _, _ in STATUS_PALETTE}
    items = []
    for z in records:
        code, _ = operational_status(z['result'])
        counts[code] += 1
        items.append((z['result']['settings']['metric'], code, z['result']['summary'],
                      z['result'].get('spec_defined', False)))
    slide = base(prs, 'Resumen de cumplimiento', heading,
                 subtitle or 'Estado por parámetro según especificaciones, control I–MR y datos omitidos.')
    for i, (name, accent, pale) in enumerate(STATUS_PALETTE):
        x = .54 + i * 3.12
        shape(slide, x, 1.91, 3.00, 1.15, pale, accent)
        text(slide, name, x+.14, 2.07, 2.72, .22, 11, accent, True, PP_ALIGN.CENTER)
        text(slide, str(counts[name]), x+.13, 2.39, 2.74, .45, 26, accent, True, PP_ALIGN.CENTER)
    text(slide, 'DESGLOSE DE PARÁMETROS', .57, 3.34, 9, .26, 10.5, TEAL, True)
    shape(slide, .54, 3.70, 12.27, .44, NAVY, radius=False)
    text(slide, 'Parámetro / medición', .77, 3.80, 5.9, .2, 10, WHITE, True)
    text(slide, 'Lecturas', 7.05, 3.80, 1.24, .2, 10, WHITE, True)
    text(slide, 'Fuera', 8.63, 3.80, 1.13, .2, 10, WHITE, True)
    text(slide, 'Estado', 10.31, 3.80, 2.09, .2, 10, WHITE, True)
    for i, (name, code, summary, spec) in enumerate(items[:9]):
        y = 4.19 + i * .258
        if i % 2 == 0:
            shape(slide, .54, y-.02, 12.27, .26, PALE, radius=False)
        col = next(accent for label, accent, _ in STATUS_PALETTE if label == code)
        text(slide, name[:61], .75, y+.02, 6.15, .21, 9.4, INK)
        text(slide, str(summary['total']), 7.08, y+.02, 1.28, .21, 9.4, INK)
        text(slide, str(summary['below']+summary['above']) if spec else 'N/D',
             8.65, y+.02, 1.19, .21, 9.4, INK)
        text(slide, code, 10.31, y+.02, 2.1, .21, 9.2, col, True)
    if len(items) > 9:
        text(slide, f'... y {len(items)-9} parámetros adicionales detallados en sus cartas.',
             .75, 6.66, 11.6, .2, 9, MUTED)
    text(slide, 'Criterios: rojo = excedencias; amarillo = alertas estadísticas / omisiones; verde = sin alertas detectadas; no evaluable = sin especificaciones.',
         .65, 6.91, 12.04, .21, 8.4, MUTED)
    return slide


def add_file_divider(prs, info, records, number):
    slide = base(prs, f'Archivo archivado {number:02d}', info['filename'][:66],
                 f'Fecha de registro: {info.get("created_at", "")[:19].replace("T", "  ")} · {info.get("kind", "Análisis")}')
    n = sum(z['result']['summary']['total'] for z in records)
    counts = [sum(operational_status(z['result'])[0] == status for z in records)
              for status, _, _ in STATUS_PALETTE]
    add_kpi(slide, .69, 1.89, 3.76, 'Parámetros analizados', str(len(records)), TEAL)
    add_kpi(slide, 4.79, 1.89, 3.75, 'Mediciones procesadas', str(n), NAVY)
    add_kpi(slide, 8.87, 1.89, 3.74, 'Alertas rojas', str(counts[2]), RED if counts[2] else GREEN)
    text(slide, 'SEMÁFORO DEL ARCHIVO', .77, 3.47, 11, .28, 11, TEAL, True)
    for i, (status, color, tint) in enumerate(STATUS_PALETTE):
        x = .72+i*3.09
        shape(slide, x, 3.96, 2.97, .98, tint, color)
        text(slide, status, x+.09, 4.09, 2.79, .23, 10.2, color, True, PP_ALIGN.CENTER)
        text(slide, str(counts[i]), x+.11, 4.43, 2.74, .27, 17, color, True, PP_ALIGN.CENTER)
    shape(slide, .7, 5.72, 11.98, .71, PALE, LINE)
    text(slide, 'A continuación: cartas individuales, distribuciones, observaciones y lista de incidencias de este archivo.',
         .92, 5.92, 11.52, .27, 11, INK)
    return slide


def add_method(prs):
    slide=base(prs,'Metodología','Cómo se interpretan estas cartas',
               'Se conservan por separado la aceptación por especificación y la estabilidad estadística.')
    rows=[
        ('01','Carta de individuales (I)', 'Línea central = media. Control: media ± 3 × (MR promedio / 1.128).'),
        ('02','Carta de rango móvil (MR)','Rango entre lecturas consecutivas. Límite superior: 3.267 × MR promedio.'),
        ('03','Capacidad del proceso', 'Cpk = distancia al límite más cercano / (3 × sigma dentro); sigma dentro = MR promedio / 1.128.'),
        ('04','Normalidad y distribución','Histograma y gráfico Q–Q. Shapiro–Wilk cuando aplica; curva normal es descriptiva.'),
        ('05','Confianza de la media', 'Intervalo bilateral al 90 % con t de Student; no sustituye límites contractuales.'),
    ]
    for i,(num,title,body) in enumerate(rows):
        y=1.63+i*.91
        shape(slide,.55,y,12.15,.77,PALE,LINE)
        shape(slide,.69,y+.14,.45,.45,TEAL)
        text(slide,num,.76,y+.255,.32,.15,9,WHITE,True,PP_ALIGN.CENTER)
        text(slide,title,1.34,y+.095,10.9,.25,13,NAVY,True)
        text(slide,body,1.34,y+.38,10.9,.28,10.5,MUTED)
    text(slide,'Semáforo Calibra: ROJO = existe incumplimiento; AMARILLO = sin excedencias, con señales de control u omisiones; VERDE = sin incidencias detectadas; NO EVALUABLE = faltan límites.',.66,6.35,11.9,.55,11,INK)
    text(slide,'La tolerancia 90/90 de los ejemplos no se calcula en esta versión; no se atribuyen resultados no verificados.',.67,6.96,11.7,.2,8.4,AMBER)


def add_metric_overview(prs, record, number):
    r=record['result'];s=r['summary'];settings=r['settings'];a=advanced(r);code,color=operational_status(r)
    slide=base(prs,f'Análisis {number:02d} / {settings["sheet"]}',settings['metric'][:72],
               f'Fuente: {record["filename"]}  ·  Especificación: {limitation(r)}' +
               (f'  ·  {period_label(r)}' if period_label(r) else ''))
    metrics=[('Muestras',str(s['total']),NAVY),('Media',fmt(s['mean'],3),TEAL),
             ('Cumplimiento',f'{fmt(s["percent"],1)}%' if r['spec_defined'] else 'N/D',GREEN if s['within']==s['total'] and r['spec_defined'] else AMBER),
             ('Cpk',fmt(a['cpk'],3),NAVY)]
    for i,(label,val,col) in enumerate(metrics):add_kpi(slide,.55+i*3.08,1.62,2.85,label,val,col)
    shape(slide,.55,2.85,8.05,2.75,WHITE,LINE)
    text(slide,'CONCLUSIÓN DEL ANÁLISIS',.77,3.04,7,.26,11,TEAL,True)
    if not r['spec_defined']: msg='No hay límites de especificación configurados. Se puede revisar estabilidad, pero no aprobar cumplimiento.'
    elif s['above']+s['below']:
        msg=f'{s["above"]+s["below"]} de {s["total"]} mediciones quedan fuera de especificación: {s["below"]} por debajo y {s["above"]} por encima.'
    else: msg=f'Las {s["total"]} mediciones válidas cumplen los límites configurados.'
    text(slide,msg,.78,3.48,7.58,.72,17,NAVY,True)
    secondary=f'Señales fuera de límites I: {s["spc_violations"]}   |   MR: {s["mr_violations"]}   |   Omitidas: {s["omitted"]}'
    text(slide,secondary,.79,4.58,7.56,.42,11,MUTED)
    shape(slide,8.89,2.85,3.88,2.75,PALE,LINE)
    text(slide,'SEMÁFORO OPERATIVO',9.13,3.12,3.35,.25,10,MUTED,True)
    shape(slide,9.13,3.66,3.13,.73,color)
    text(slide,code,9.17,3.85,3.02,.26,16,WHITE,True,PP_ALIGN.CENTER)
    text(slide,'Fuera especificación ≠ fuera de control estadístico',9.14,4.74,3.29,.47,10,INK)
    shape(slide,.56,5.89,12.2,.92,PALE,LINE)
    normal=f'Shapiro p = {fmt(a["shapiro"],3)}' if a['shapiro'] is not None else 'Shapiro: no evaluable'
    ci=f'IC 90 % para media: [{fmt(a["ci"][0],3)}; {fmt(a["ci"][1],3)}]'
    text(slide,normal + '     ·     ' + ci + '     ·     Ppk = ' + fmt(a['ppk'],3),.8,6.18,11.5,.36,10.5,NAVY)


def add_metric_charts(prs, record, number):
    r=record['result'];metric=r['settings']['metric']
    slide=base(prs,f'Cartas y distribución / {number:02d}',f'{metric[:63]} — diagnóstico visual',
               'Datos reales del análisis · especificaciones en rojo · límites de control en azul punteado.')
    positions=[(.52,1.65,6.03,2.48,'I','Individuales X — control y especificación'),
               (6.75,1.65,6.03,2.48,'H','Histograma y curva normal'),
               (.52,4.34,6.03,2.48,'MR','Rangos móviles (mR)'),
               (6.75,4.34,6.03,2.48,'QQ','Gráfico normal Q–Q')]
    for x,y,w,h,kind,label in positions:
        plot_card(slide,r,kind,x,y,w,h,label)


def add_issues(prs, record, number):
    r=record['result'];s=r['summary'];l=r['limits']
    issues=[]
    for p in r['points']:
        types=[]
        if not p['in_specs'] and r['spec_defined']: types.append('Fuera especificación')
        if not p['in_control']: types.append('Fuera control I')
        if p['mr'] is not None and p['mr']>l['mr_ucl']: types.append('Salto MR')
        if types:
            issues.append((p['index'],p['row'],p['label'],fmt(p['value'],4),' / '.join(types)))
    for o in r.get('omitted',[]):
        issues.append(('—',o['row'],o['label'],str(o.get('value',''))[:19],'Lectura omitida'))
    if not issues: return
    per_page=17
    for page in range(0,len(issues),per_page):
        slide=base(prs,f'Incidencias / análisis {number:02d}',r['settings']['metric'][:68],
                   f'Listado completo de incidencias detectadas · {len(issues)} filas en total · página {page//per_page+1} de {math.ceil(len(issues)/per_page)}')
        shape(slide,.56,1.69,12.24,.45,NAVY,radius=False)
        cx=[.73,1.52,2.25,6.75,9.0]
        headers=['N.º','Fila Excel','Etiqueta','Valor','Observación']
        widths=[.7,.9,4.38,1.3,3.47]
        for xx,wd,hh in zip(cx,widths,headers):text(slide,hh,xx,1.78,wd,.21,10,WHITE,True)
        for i,item in enumerate(issues[page:page+per_page]):
            y=2.18+i*.269
            if i%2==0:shape(slide,.56,y-.03,12.24,.267,PALE,radius=False)
            for j,(xx,wd,v) in enumerate(zip(cx,widths,item)):
                text(slide,v,xx,y,wd,.19,8.8,RED if j==4 and 'Fuera especificación' in str(v) else INK,j==4)


def add_conclusions(prs, records):
    slide=base(prs,'Cierre','Conclusiones y próximos pasos',
               'Síntesis basada exclusivamente en las mediciones y los límites guardados en cada análisis.')
    red=[z for z in records if operational_status(z['result'])[0]=='ROJO']
    yellow=[z for z in records if operational_status(z['result'])[0]=='AMARILLO']
    blank=[z for z in records if operational_status(z['result'])[0]=='NO EVALUABLE']
    notes=[
        (RED,'Revisar incumplimientos',f'{len(red)} análisis con una o más mediciones fuera de especificación.'),
        (AMBER,'Investigar señales estadísticas',f'{len(yellow)} análisis sin incumplimientos, pero con alertas de control u omisiones.'),
        (MUTED,'Confirmar especificaciones',f'{len(blank)} análisis sin límites configurados; su cumplimiento queda NO EVALUABLE.'),
        (TEAL,'Conservar trazabilidad','Los análisis y las filas de origen quedan identificados; el Excel original no se modifica.'),
    ]
    for i,(accent,title,body) in enumerate(notes):
        y=1.72+i*1.17
        shape(slide,.56,y,12.17,.99,PALE,LINE)
        shape(slide,.67,y+.19,.065,.61,accent,radius=False)
        text(slide,title,.96,y+.12,11,.32,16,NAVY,True)
        text(slide,body,.97,y+.53,11,.23,11,MUTED)
    text(slide,'Este informe usa reglas I–MR y verificación de especificaciones. No constituye certificación contractual ni sustituye la revisión de ingeniería.',.67,6.65,11.8,.44,10,INK)


def create_pptx(records, title='Reporte de control estadístico', groups=None):
    """PowerPoint por análisis, por archivo o compilado de archivos independientes."""
    records = [normalize_record(z) for z in records]
    if not records:
        raise ValueError('No hay análisis guardados para exportar.')
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    prs.core_properties.title=title
    prs.core_properties.subject='Cartas X–mR y resultados de control de calidad'
    prs.core_properties.author='Calibra — Informe automático'
    add_title(prs,title,records)
    add_dashboard(prs,records)
    add_semaforo(prs,records)
    add_method(prs)
    if groups:
        count=0
        for group_number, (metadata, subset) in enumerate(groups,1):
            entries=[normalize_record(z) for z in subset]
            add_file_divider(prs,metadata,entries,group_number)
            add_semaforo(prs, entries, heading='Semáforo · '+metadata['filename'][:44],
                         subtitle='Desglose exclusivo de las mediciones de este archivo.')
            for record in entries:
                count+=1
                add_metric_overview(prs,record,count)
                add_metric_charts(prs,record,count)
                add_issues(prs,record,count)
    else:
        for ix,record in enumerate(records,1):
            add_metric_overview(prs,record,ix)
            add_metric_charts(prs,record,ix)
            add_issues(prs,record,ix)
    add_conclusions(prs,records)
    buffer=io.BytesIO();prs.save(buffer);buffer.seek(0)
    return buffer
